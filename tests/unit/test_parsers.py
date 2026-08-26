"""Unit tests for document parsers and MIME sniffing (no database required)."""

import io

import pytest
from docx import Document as DocxBuilder
from reportlab.lib.pagesizes import LETTER
from reportlab.pdfgen import canvas

from app.core.exceptions import DocumentParseError, UnsupportedDocumentError
from app.ingestion.parsers.csv import parse_csv
from app.ingestion.parsers.docx import parse_docx
from app.ingestion.parsers.pdf import parse_pdf
from app.ingestion.parsers.registry import (
    CSV_MIME,
    DOCX_MIME,
    MARKDOWN_MIME,
    PDF_MIME,
    TEXT_MIME,
    get_parser,
    sniff_mime_type,
)
from app.ingestion.parsers.text import parse_markdown, parse_text


def make_pdf(pages: list[list[str]], *, heading: str | None = None) -> bytes:
    buffer = io.BytesIO()
    pdf = canvas.Canvas(buffer, pagesize=LETTER)
    for page_index, lines in enumerate(pages):
        y = 750.0
        if heading is not None and page_index == 0:
            pdf.setFont("Helvetica-Bold", 18)
            pdf.drawString(72, y, heading)
            y -= 50
        pdf.setFont("Helvetica", 10)
        for line in lines:
            pdf.drawString(72, y, line)
            y -= 14
        pdf.showPage()
    pdf.save()
    return buffer.getvalue()


def make_docx() -> bytes:
    document = DocxBuilder()
    document.add_heading("Quarterly Report", level=1)
    document.add_paragraph("The first quarter went well. Revenue grew steadily.")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Item"
    table.cell(0, 1).text = "Amount"
    table.cell(1, 0).text = "Widgets"
    table.cell(1, 1).text = "1200"
    document.add_paragraph("Closing remarks follow the table.")
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


class TestPdfParser:
    def test_pages_and_text(self):
        data = make_pdf([["Alpha line one.", "Alpha line two."], ["Beta line one."]])
        output = parse_pdf(data)
        assert output.page_count == 2
        assert {element.page_number for element in output.elements} == {1, 2}
        all_text = "\n".join(element.text for element in output.elements)
        assert "Alpha line one." in all_text
        assert "Beta line one." in all_text

    def test_heading_detected(self):
        data = make_pdf(
            [["Body text sentence one.", "Body text sentence two."]],
            heading="Invoice Summary",
        )
        output = parse_pdf(data)
        headings = [e for e in output.elements if e.element_type == "heading"]
        assert any(h.text == "Invoice Summary" for h in headings)
        assert any(e.element_type == "paragraph" for e in output.elements)

    def test_corrupt_pdf_raises(self):
        with pytest.raises(DocumentParseError):
            parse_pdf(b"%PDF-1.7\nthis is not really a pdf")


class TestDocxParser:
    def test_structure_in_order(self):
        output = parse_docx(make_docx())
        types = [element.element_type for element in output.elements]
        assert types == ["heading", "paragraph", "table", "paragraph"]
        assert output.elements[0].text == "Quarterly Report"
        table = output.elements[2]
        assert "| Item | Amount |" in table.text
        assert "| Widgets | 1200 |" in table.text
        assert table.metadata["rows"] == 2
        assert table.metadata["columns"] == 2

    def test_corrupt_docx_raises(self):
        with pytest.raises(DocumentParseError):
            parse_docx(b"PK\x03\x04 garbage [Content_Types].xml")


class TestTextParser:
    def test_txt_paragraphs(self):
        data = b"First paragraph line.\n\nSecond paragraph line.\n\n\nThird one."
        output = parse_text(data)
        assert [e.element_type for e in output.elements] == ["paragraph"] * 3
        assert output.elements[1].text == "Second paragraph line."
        assert output.page_count is None

    def test_markdown(self):
        md = "\n".join(
            [
                "# Title Heading",
                "",
                "Intro paragraph text.",
                "",
                "## Details",
                "",
                "- first item",
                "- second item",
                "",
                "```",
                "x = 1",
                "```",
            ]
        ).encode()
        output = parse_markdown(md)
        types = [e.element_type for e in output.elements]
        assert types == ["heading", "paragraph", "heading", "list", "other"]
        assert output.elements[0].text == "Title Heading"
        assert output.elements[0].metadata["level"] == 1
        assert output.elements[2].text == "Details"
        assert output.elements[2].metadata["level"] == 2
        assert "first item" in output.elements[3].text
        assert "x = 1" in output.elements[4].text


class TestCsvParser:
    def test_slices_and_metadata(self):
        rows = "\n".join(f"item{i},{i}" for i in range(65))
        data = f"name,value\n{rows}\n".encode()
        output = parse_csv(data)
        assert [e.element_type for e in output.elements] == ["table"] * 3
        assert output.elements[0].metadata["columns"] == ["name", "value"]
        assert output.elements[0].metadata["row_count"] == 30
        assert output.elements[1].metadata["row_start"] == 30
        assert output.elements[2].metadata["row_count"] == 5
        assert "| item0 | 0 |" in output.elements[0].text
        assert "| item64 | 64 |" in output.elements[2].text

    def test_empty_bytes_raise(self):
        with pytest.raises(DocumentParseError):
            parse_csv(b"")


class TestMimeSniffing:
    def test_pdf_magic_wins_over_extension(self):
        assert sniff_mime_type(b"%PDF-1.4 rest of file", "notes.txt") == PDF_MIME

    def test_real_docx_detected_regardless_of_extension(self):
        assert sniff_mime_type(make_docx(), "report.bin") == DOCX_MIME

    def test_zip_magic_with_docx_extension(self):
        assert sniff_mime_type(b"PK\x03\x04somezipdata", "file.docx") == DOCX_MIME

    def test_zip_that_is_not_docx_rejected(self):
        with pytest.raises(UnsupportedDocumentError):
            sniff_mime_type(b"PK\x03\x04somezipdata", "archive.zip")

    def test_text_extensions(self):
        assert sniff_mime_type(b"hello world", "a.txt") == TEXT_MIME
        assert sniff_mime_type(b"# hi", "a.md") == MARKDOWN_MIME
        assert sniff_mime_type(b"a,b\n1,2", "a.csv") == CSV_MIME

    def test_fake_pdf_extension_not_trusted(self):
        with pytest.raises(UnsupportedDocumentError):
            sniff_mime_type(b"just plain text", "fake.pdf")

    def test_binary_garbage_rejected(self):
        with pytest.raises(UnsupportedDocumentError):
            sniff_mime_type(b"\x00\xff\xfe\x01binary", "data.bin")

    def test_decodable_unknown_extension_rejected(self):
        with pytest.raises(UnsupportedDocumentError):
            sniff_mime_type(b"plain text content", "script.py")

    def test_get_parser_unknown_mime(self):
        with pytest.raises(UnsupportedDocumentError):
            get_parser("application/zip")
