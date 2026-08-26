"""End-to-end pipeline tests: bytes in, chunks out (no database required)."""

import hashlib
import io

import pytest
from docx import Document as DocxBuilder
from reportlab.lib.pagesizes import LETTER
from reportlab.pdfgen import canvas

from app.core.exceptions import DocumentParseError, UnsupportedDocumentError
from app.ingestion.parsers.registry import CSV_MIME, DOCX_MIME, PDF_MIME, TEXT_MIME
from app.ingestion.pipeline import PARSER_VERSION, run_pipeline


def make_pdf(pages: list[list[str]]) -> bytes:
    buffer = io.BytesIO()
    pdf = canvas.Canvas(buffer, pagesize=LETTER)
    for lines in pages:
        y = 750.0
        pdf.setFont("Helvetica", 10)
        for line in lines:
            pdf.drawString(72, y, line)
            y -= 14
        pdf.showPage()
    pdf.save()
    return buffer.getvalue()


def make_docx() -> bytes:
    document = DocxBuilder()
    document.add_heading("Pipeline Fixture", level=1)
    document.add_paragraph("A paragraph that flows through the whole pipeline.")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Key"
    table.cell(0, 1).text = "Value"
    table.cell(1, 0).text = "Total"
    table.cell(1, 1).text = "42"
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def test_pdf_pipeline():
    data = make_pdf([["Page one sentence."], ["Page two sentence."]])
    result = run_pipeline(data, "report.pdf")
    assert result.mime_type == PDF_MIME
    assert result.page_count == 2
    assert result.size_bytes == len(data)
    assert result.sha256 == hashlib.sha256(data).hexdigest()
    assert result.parser_version == PARSER_VERSION == "1.0.0"
    assert result.elements
    assert result.chunks
    assert [chunk.order_index for chunk in result.chunks] == list(range(len(result.chunks)))


def test_txt_pipeline_sha_stable():
    data = b"Stable text paragraph one.\n\nStable text paragraph two."
    first = run_pipeline(data, "notes.txt")
    second = run_pipeline(data, "notes.txt")
    assert first.mime_type == TEXT_MIME
    assert first.page_count is None
    assert first.sha256 == second.sha256 == hashlib.sha256(data).hexdigest()
    assert first.chunks == second.chunks
    assert "Stable text paragraph one." in first.chunks[0].text


def test_docx_pipeline():
    data = make_docx()
    result = run_pipeline(data, "fixture.docx")
    assert result.mime_type == DOCX_MIME
    assert result.sha256 == hashlib.sha256(data).hexdigest()
    assert any(element.element_type == "table" for element in result.elements)
    assert any("| Key | Value |" in chunk.text for chunk in result.chunks)
    # The heading becomes the section of its chunk.
    assert result.chunks[0].section == "Pipeline Fixture"


def test_csv_pipeline():
    rows = "\n".join(f"row{i},{i * 2}" for i in range(40))
    data = f"label,amount\n{rows}\n".encode()
    result = run_pipeline(data, "table.csv")
    assert result.mime_type == CSV_MIME
    assert result.sha256 == hashlib.sha256(data).hexdigest()
    assert all(element.element_type == "table" for element in result.elements)
    assert any("| row0 | 0 |" in chunk.text for chunk in result.chunks)


def test_pipeline_rejects_unknown_type():
    with pytest.raises(UnsupportedDocumentError):
        run_pipeline(b"\x00\xff\xfe\x01", "mystery.bin")


def test_pipeline_propagates_parse_error():
    with pytest.raises(DocumentParseError):
        run_pipeline(b"%PDF-1.7\nbroken beyond repair", "broken.pdf")
